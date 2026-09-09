import sys, os, re, json, time, urllib.request
from datetime import datetime, timezone, timedelta
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

SHEET_ID = '1EQ86CzyKT0JX__siFIZW0u9ZuAmk25iudGXXlX2xcbk'
SHEET_URL = f'https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?usp=sharing'
EXPORT_URL = f'https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx'

def download_sheet(target_file='temp_sheet.xlsx'):
    print(f'[1/4] Downloading latest Google Sheet (ID: {SHEET_ID})...')
    req = urllib.request.Request(EXPORT_URL, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=120) as resp, open(target_file, 'wb') as f:
        while True:
            chunk = resp.read(1024 * 1024)
            if not chunk: break
            f.write(chunk)
    print('      Download complete.')
    return target_file

def extract_data(excel_file):
    print('[2/4] Extracting data from sheets...')
    wb = openpyxl.load_workbook(excel_file, data_only=True, read_only=True)
    overall = []
    ws_sep = wb['สรุปแยก']
    for row in ws_sep.iter_rows(min_row=5, max_row=34, min_col=1, max_col=6, values_only=True):
        no = int(float(row[0])) if row[0] is not None else len(overall) + 1
        name = str(row[1] or '').strip()
        room = str(row[2] or '').strip()
        lineup = str(row[3] or '-').strip()
        hr = str(row[4] or '-').strip()
        cs = str(row[5] or '-').strip()
        m_hr = re.search(r'จำนวน\s*(\d+)\s*ครั้ง', hr)
        hr_cnt = int(m_hr.group(1)) if m_hr else 0
        m_cs = re.search(r'จำนวน\s*(\d+)\s*ครั้ง', cs)
        cs_cnt = int(m_cs.group(1)) if m_cs else 0
        overall.append({
            'no': no, 'name': name, 'room': room,
            'lineup_summary': lineup,
            'homeroom_summary': hr, 'chitsuksa_summary': cs,
            'homeroom_count': hr_cnt, 'chitsuksa_count': cs_cnt,
            'total_missed': hr_cnt + cs_cnt
        })

    monthly_sheets = {
        '5': 'เข้าแถว พ.ค จพศ', '6': 'เข้าแถว มิ.ย จพศ',
        '7': 'เข้าแถว ก.ค จพศ', '8': 'เข้าแถว ส.ค. จพศ',
        '9': 'เข้าแถว ก.ย. จพศ'
    }
    monthly = {}
    for m_key, s_name in monthly_sheets.items():
        ws_m = wb[s_name]
        m_list = []
        for row in ws_m.iter_rows(min_row=5, max_row=34, min_col=1, max_col=5, values_only=True):
            no = int(float(row[0])) if row[0] is not None else len(m_list) + 1
            name = str(row[1] or '').strip()
            room = str(row[2] or '').strip()
            c4, c5 = row[3], row[4]
            missed = 0
            cnt_str = '-'
            if c4 is not None and str(c4).strip() not in ['', '-', '0', '0.0']:
                s = str(c4).strip()
                try:
                    missed = int(float(s))
                    cnt_str = f'{missed} ครั้ง'
                except ValueError:
                    cnt_str = s
                    m = re.search(r'(\d+)', s)
                    if m: missed = int(m.group(1))
            reas_str = str(c5 or '-').strip()
            if missed == 0 or cnt_str == '-':
                sum_str, reas_str, cnt_str = '-', '-', '-'
            else:
                sum_str = f'ไม่เข้าแถว จำนวน {cnt_str}\n{reas_str}' if reas_str != '-' else f'ไม่เข้าแถว จำนวน {cnt_str}'
            m_list.append({
                'no': no, 'name': name, 'room': room,
                'missed_count': missed,
                'count_display': cnt_str,
                'reason_display': reas_str,
                'summary_display': sum_str
            })
        monthly[m_key] = m_list
    wb.close()
    return overall, monthly

def get_thai_time():
    tz = timezone(timedelta(hours=7))
    now = datetime.now(tz)
    all_th = ['', 'ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.', 'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.']
    return f'{now.day} {all_th[now.month]} {now.year+543} เวลา {now.strftime("%H:%M:%S")} น.'

def update_index(overall, monthly, html_path='index.html'):
    print(f'[3/4] Updating {html_path}...')
    with open(html_path, 'r', encoding='utf-8') as f:
        html = f.read()
    oj = json.dumps(overall, ensure_ascii=False)
    mj = json.dumps(monthly, ensure_ascii=False)
    html = re.sub(r'const overallData = \[.*?\];', f'const overallData = {oj};', html, flags=re.DOTALL)
    html = re.sub(r'const monthlyData = \{.*?\};', f'const monthlyData = {mj};', html, flags=re.DOTALL)
    t_str = get_thai_time()
    
    badge_style = 'display:inline-flex;align-items:center;gap:6px;background:rgba(16,185,129,0.25);border:1px solid rgba(16,185,129,0.4);padding:5px 12px;border-radius:9999px;color:#d1fae5;font-weight:500;'
    new_badge = f'<span id="updateTimeBadge" style="{badge_style}">⏰ อัพเดทล่าสุด: {t_str}</span>'
    
    if '<div class="meta-badges"' in html:
        html = re.sub(r'<span id="updateTimeBadge"[^>]*>.*?</span>', new_badge, html)
    else:
        link_style = 'display:inline-flex;align-items:center;gap:6px;background:rgba(255,255,255,0.2);backdrop-filter:blur(8px);padding:5px 12px;border-radius:9999px;color:#fff;text-decoration:none;font-weight:500;border:1px solid rgba(255,255,255,0.3);'
        meta_html = f'<div class="meta-badges" style="display:flex;flex-wrap:wrap;gap:10px;align-items:center;margin-top:14px;font-size:13px;"><a href="{SHEET_URL}" target="_blank" style="{link_style}"><span>🔗 ลิงก์ชีทออนไลน์ (ID: {SHEET_ID[:12]}...)</span> ↗</a>{new_badge}</div>'
        html = html.replace('<div class="legend-banner">', meta_html + '\n        <div class="legend-banner">')
        
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f'[4/4] index.html updated successfully! ({t_str})')

def main():
    t0 = time.time()
    target = 'temp_download.xlsx'
    if len(sys.argv) > 1 and os.path.exists(sys.argv[1]):
        efile = sys.argv[1]
    elif os.path.exists('latest_online.xlsx'):
        efile = 'latest_online.xlsx'
    else:
        efile = download_sheet(target)
    ov, mo = extract_data(efile)
    update_index(ov, mo, 'index.html')
    if efile == target and os.path.exists(target):
        try: os.remove(target)
        except: pass
    print(f'Done in {time.time() - t0:.2f}s')

if __name__ == '__main__':
    main()
